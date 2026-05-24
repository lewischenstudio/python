# from pdfminer.high_level import extract_text

# text = extract_text("12345.pdf")
# print(text)

from docx import Document

# # Open an existing document
# document = Document("plan.docx")

# # Add a new paragraph at the end of the document
# document.add_paragraph("This is newly appended text.")

# # Save the modified document
# document.save("updated_file.docx")

# document = Document("updated_file.docx")

# signature = "LC"
# for paragraph in document.paragraphs:
#     print(paragraph.text)

#     if "{{INSERT_HERE}}" in paragraph.text:
#         paragraph.text = f"This is newly appended text.   {signature}"

# # # Save the modified document
# document.save("updated_file.docx")

# pip install pymupdf

import pymupdf  # This imports the library


def convert_pdf_to_png_pymupdf(pdf_path, dpi=300):
    doc = pymupdf.open(pdf_path)
    for i, page in enumerate(doc):
        # Render the page to a Pixmap with a specified DPI for quality control
        pix = page.get_pixmap(dpi=dpi)
        output_image_path = f"page-{i+1}.png"

        # Save the pixmap as a PNG file
        pix.save(output_image_path)
        print(f"Saved {output_image_path}")
    doc.close()


# Example usage
pdf_file = "abc.pdf"
convert_pdf_to_png_pymupdf(pdf_file)
